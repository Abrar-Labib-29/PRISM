"""
iValue PRISM — Proposal & Bill of Materials (BOM) Exporters
SRS References: §3 Phase 2 step 3, §8.3, §8.4, §10.5, §9.10
Implementation Plan: TASK-P1.9

This module handles:
  1. Internal dataclasses: BomItem, BomExportRequest, BoqExportRequest, ExportResponse (§8.3).
  2. BOMExporter: Generates formatted Excel BOM (.xlsx) with strictly blank pricing columns (§8.4).
     - 14 standard columns
     - Prominent red watermark/disclaimer: "DRAFT — PRICING NOT INCLUDED — FOR INTERNAL USE ONLY"
     - "To be filled by Sales" headers for pricing
     - Zero-tolerance: no prices, costs, fees, or monetary formulas anywhere
     - Robust fallback to CSV on openpyxl rendering failure (§10.5)
     - Disk-full and OS error guards (§10.5)
"""

import csv
from dataclasses import dataclass, field
from datetime import datetime
import logging
import os
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    import docx
    from docx.enum.section import WD_ORIENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor
except ImportError:
    docx = None

_logger = logging.getLogger("prism.exporters")

# Standard Disclaimer & Notices per §8.4 & §10.5
BOM_DISCLAIMER_TEXT = "DRAFT — PRICING NOT INCLUDED — FOR INTERNAL USE ONLY"
BOM_COMPANY_HEADER = "iValue InfoSolutions — Bill of Materials"
SALES_NOTICE_TEXT = "Pricing to be completed by iValue Sales Team. Contact: sales@ivalue.co.in"
TBD_SALES_FALLBACK = "TBD — Consult Sales"

# The 14 Canonical Column Headers per §8.4
BOM_HEADERS: List[str] = [
    "S.No.",
    "OEM",
    "Product Name",
    "Product Category",
    "Model/Part Reference",
    "Quantity",
    "Licensing Model",
    "License Term",
    "Licensing Unit",
    "Deployment Model",
    "Support Tier",
    "Unit Price — To be filled by Sales",
    "Total Price — To be filled by Sales",
    "Remarks",
]


@dataclass
class BomItem:
    """Line item in a Bill of Materials or Bill of Quantities (§8.3)."""
    product_id: str
    quantity: int = 1
    license_term: str = "1-Year"
    support_tier: str = "Standard"
    remarks: str = ""
    # Optional overrides from grid / caller
    oem: Optional[str] = None
    product_name: Optional[str] = None
    product_category: Optional[str] = None
    licensing_model: Optional[str] = None
    licensing_unit: Optional[str] = None
    deployment_model: Optional[str] = None


@dataclass
class BomExportRequest:
    """Request contract for Excel BOM export (§8.3)."""
    query_id: str
    customer_name: str
    items: List[BomItem]
    output_path: str                     # Local destination path for .xlsx


@dataclass
class BoqExportRequest:
    """Request contract for Word BOQ export (§8.3)."""
    query_id: str
    customer_name: str
    project_reference: str
    prepared_by: str
    items: List[BomItem]
    output_path: str                     # Local destination path for .docx
    currency_code: str = ""              # Optional currency indicator (default: neutral blank)


@dataclass
class ExportResponse:
    """Result of document generation operations (§8.3)."""
    status: str                          # "success" | "error"
    file_path: str                       # Resolved absolute local file path
    error_message: Optional[str] = None


class BOMExporter:
    """
    Generates formatted Excel BOM (.xlsx) with blank pricing columns per §8.4.
    """

    def __init__(self, catalog: Optional[Dict[str, Dict[str, str]]] = None) -> None:
        """
        catalog: Optional pre-loaded mapping of product_id -> product attributes.
        If omitted, loads automatically from dataset assets.
        """
        self._catalog: Dict[str, Dict[str, str]] = catalog if catalog is not None else {}
        if not self._catalog:
            self._load_catalog()

    def _load_catalog(self) -> None:
        """Loads product attributes from raw dataset or json assets."""
        dataset_path = os.path.join(
            os.path.dirname(__file__), "..", "..", "data", "raw", "iValue_Solution_Recommendation_Dataset.xlsx"
        )
        dataset_path = os.path.abspath(dataset_path)

        if os.path.isfile(dataset_path):
            try:
                wb = openpyxl.load_workbook(dataset_path, data_only=True, read_only=True)
                if "Products" in wb.sheetnames:
                    p_sheet = wb["Products"]
                    for row in p_sheet.iter_rows(min_row=2, values_only=True):
                        if row and row[0]:
                            pid = str(row[0]).strip()
                            self._catalog[pid] = {
                                "oem_name": str(row[2]).strip() if len(row) > 2 and row[2] else "",
                                "product_category": str(row[5]).strip() if len(row) > 5 and row[5] else "",
                                "product_name": str(row[6]).strip() if len(row) > 6 and row[6] else "",
                                "deployment_model": str(row[11]).strip() if len(row) > 11 and row[11] else "",
                                "licensing_model": "",
                                "licensing_unit": "",
                            }
                if "Product_Commercial" in wb.sheetnames:
                    c_sheet = wb["Product_Commercial"]
                    for row in c_sheet.iter_rows(min_row=2, values_only=True):
                        if row and row[0]:
                            pid = str(row[0]).strip()
                            if pid in self._catalog:
                                self._catalog[pid]["licensing_model"] = (
                                    str(row[3]).strip() if len(row) > 3 and row[3] else ""
                                )
                                self._catalog[pid]["licensing_unit"] = (
                                    str(row[5]).strip() if len(row) > 5 and row[5] else ""
                                )
                return
            except Exception as e:
                _logger.warning(f"Could not load Excel dataset catalog for BOM: {e}")

        # Fallback to composite_products.json
        json_path = os.path.join(
            os.path.dirname(__file__), "..", "..", "data", "composite_products.json"
        )
        json_path = os.path.abspath(json_path)
        if os.path.isfile(json_path):
            try:
                import json
                with open(json_path, "r", encoding="utf-8") as f:
                    products = json.load(f)
                for item in products:
                    pid = item.get("product_id") or ""
                    if pid:
                        self._catalog[pid] = {
                            "oem_name": item.get("oem_name") or "",
                            "product_category": item.get("product_category") or "",
                            "product_name": item.get("product_name") or "",
                            "deployment_model": item.get("deployment_model") or "",
                            "licensing_model": item.get("licensing_model") or "",
                            "licensing_unit": item.get("licensing_unit") or "",
                        }
            except Exception as e:
                _logger.warning(f"Could not load json catalog for BOM: {e}")

    def export(self, request: BomExportRequest) -> ExportResponse:
        """
        Creates .xlsx via openpyxl with columns per §8.4 BOM format:
        S.No. | OEM | Product Name | Product Category | Model/Part Reference |
        Quantity | Licensing Model | License Term | Licensing Unit |
        Deployment Model | Support Tier | Unit Price | Total Price | Remarks

        - Unit Price header: 'Unit Price — To be filled by Sales'
        - Total Price header: 'Total Price — To be filled by Sales'
        - Unit Price and Total Price cells: ALWAYS BLANK.
        - Row 1 disclaimer: 'DRAFT — PRICING NOT INCLUDED — FOR INTERNAL USE ONLY'
        - Standard iValue InfoSolutions header row.
        - Missing fields: fill with 'TBD — Consult Sales' per §10.5.
        - Column widths auto-sized for readability.
        """
        out_path = os.path.abspath(request.output_path)
        out_dir = os.path.dirname(out_path)

        # Ensure parent directory exists
        try:
            if out_dir:
                os.makedirs(out_dir, exist_ok=True)
        except OSError as e:
            _logger.critical(f"Disk or filesystem error creating directory {out_dir}: {e}")
            return ExportResponse(
                status="error",
                file_path="",
                error_message=f"Cannot create directory '{out_dir}': {e}",
            )

        try:
            wb = openpyxl.Workbook()
            ws = wb.active
            ws.title = "Bill of Materials"

            # Enable grid lines
            ws.views.sheetView[0].showGridLines = True

            # Row 1: Disclaimer (A1:N1 merged)
            ws.merge_cells("A1:N1")
            cell_disclaimer = ws["A1"]
            cell_disclaimer.value = BOM_DISCLAIMER_TEXT
            cell_disclaimer.font = Font(name="Segoe UI", size=11, bold=True, color="B91C1C")  # Dark crimson red
            cell_disclaimer.fill = PatternFill(start_color="FEF2F2", fill_type="solid")       # Subtle rose tint
            cell_disclaimer.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[1].height = 28

            # Row 2: Company Header (A2:N2 merged)
            ws.merge_cells("A2:N2")
            cell_company = ws["A2"]
            cell_company.value = BOM_COMPANY_HEADER
            cell_company.font = Font(name="Segoe UI", size=14, bold=True, color="0F172A")    # Deep Navy
            cell_company.alignment = Alignment(horizontal="center", vertical="center")
            ws.row_dimensions[2].height = 32

            # Row 3: Metadata Details
            date_str = datetime.now().strftime("%Y-%m-%d %H:%M")
            ws["A3"] = f"Customer: {request.customer_name if request.customer_name else 'Unspecified'}"
            ws["A3"].font = Font(name="Segoe UI", size=10, bold=True, color="334155")
            
            ws["G3"] = f"Date: {date_str}"
            ws["G3"].font = Font(name="Segoe UI", size=10, italic=True, color="475569")
            
            ws["K3"] = f"Query ID: {request.query_id}"
            ws["K3"].font = Font(name="Segoe UI", size=10, italic=True, color="475569")
            ws.row_dimensions[3].height = 22

            # Row 4: Column Headers
            header_row_idx = 4
            ws.row_dimensions[header_row_idx].height = 30
            header_fill = PatternFill(start_color="0F172A", fill_type="solid")  # Brand dark slate
            header_font = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
            header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

            thin_border = Border(
                left=Side(style="thin", color="CBD5E1"),
                right=Side(style="thin", color="CBD5E1"),
                top=Side(style="thin", color="CBD5E1"),
                bottom=Side(style="thin", color="CBD5E1"),
            )

            for col_idx, header_title in enumerate(BOM_HEADERS, start=1):
                cell = ws.cell(row=header_row_idx, column=col_idx, value=header_title)
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = header_align
                cell.border = thin_border

            # Data rows (Row 5+)
            align_center = Alignment(horizontal="center", vertical="center")
            align_left = Alignment(horizontal="left", vertical="center")
            data_font = Font(name="Segoe UI", size=10, color="0F172A")
            alt_fill = PatternFill(start_color="F8FAFC", fill_type="solid")
            white_fill = PatternFill(start_color="FFFFFF", fill_type="solid")

            for idx, item in enumerate(request.items, start=1):
                row_idx = header_row_idx + idx
                ws.row_dimensions[row_idx].height = 24
                cat_info = self._catalog.get(item.product_id, {})

                # Extract attributes with overrides and §10.5 fallback
                oem = item.oem or cat_info.get("oem_name") or TBD_SALES_FALLBACK
                pname = item.product_name or cat_info.get("product_name") or TBD_SALES_FALLBACK
                pcat = item.product_category or cat_info.get("product_category") or TBD_SALES_FALLBACK
                lic_model = item.licensing_model or cat_info.get("licensing_model") or TBD_SALES_FALLBACK
                lic_unit = item.licensing_unit or cat_info.get("licensing_unit") or TBD_SALES_FALLBACK
                deploy_model = item.deployment_model or cat_info.get("deployment_model") or TBD_SALES_FALLBACK

                row_fill = alt_fill if (idx % 2 == 0) else white_fill

                # 1. S.No.
                c1 = ws.cell(row=row_idx, column=1, value=idx)
                c1.alignment = align_center

                # 2. OEM
                c2 = ws.cell(row=row_idx, column=2, value=oem)
                c2.alignment = align_left

                # 3. Product Name
                c3 = ws.cell(row=row_idx, column=3, value=pname)
                c3.alignment = align_left

                # 4. Product Category
                c4 = ws.cell(row=row_idx, column=4, value=pcat)
                c4.alignment = align_left

                # 5. Model/Part Reference (Product_ID)
                c5 = ws.cell(row=row_idx, column=5, value=item.product_id)
                c5.alignment = align_center

                # 6. Quantity
                c6 = ws.cell(row=row_idx, column=6, value=item.quantity)
                c6.alignment = align_center

                # 7. Licensing Model
                c7 = ws.cell(row=row_idx, column=7, value=lic_model)
                c7.alignment = align_left

                # 8. License Term
                c8 = ws.cell(row=row_idx, column=8, value=item.license_term)
                c8.alignment = align_center

                # 9. Licensing Unit
                c9 = ws.cell(row=row_idx, column=9, value=lic_unit)
                c9.alignment = align_left

                # 10. Deployment Model
                c10 = ws.cell(row=row_idx, column=10, value=deploy_model)
                c10.alignment = align_left

                # 11. Support Tier
                c11 = ws.cell(row=row_idx, column=11, value=item.support_tier)
                c11.alignment = align_center

                # 12. Unit Price — ALWAYS BLANK
                c12 = ws.cell(row=row_idx, column=12, value=None)
                c12.alignment = align_center

                # 13. Total Price — ALWAYS BLANK
                c13 = ws.cell(row=row_idx, column=13, value=None)
                c13.alignment = align_center

                # 14. Remarks
                c14 = ws.cell(row=row_idx, column=14, value=item.remarks or "")
                c14.alignment = align_left

                # Style all data cells in this row
                for c in (c1, c2, c3, c4, c5, c6, c7, c8, c9, c10, c11, c12, c13, c14):
                    c.font = data_font
                    c.fill = row_fill
                    c.border = thin_border

            # Footer / Pricing Notice Row
            footer_row_idx = header_row_idx + len(request.items) + 1
            ws.row_dimensions[footer_row_idx].height = 24
            ws.merge_cells(start_row=footer_row_idx, start_column=1, end_row=footer_row_idx, end_column=11)
            f_cell = ws.cell(row=footer_row_idx, column=1, value=SALES_NOTICE_TEXT)
            f_cell.font = Font(name="Segoe UI", size=9, italic=True, color="64748B")
            f_cell.alignment = Alignment(horizontal="left", vertical="center")

            # Blank pricing cells on footer row
            c_f12 = ws.cell(row=footer_row_idx, column=12, value=None)
            c_f13 = ws.cell(row=footer_row_idx, column=13, value=None)
            c_f14 = ws.cell(row=footer_row_idx, column=14, value="")

            # Auto-size columns with sensible bounds
            for col in ws.columns:
                col_letter = get_column_letter(col[0].column)
                # Compute max length of non-merged cell contents
                max_len = 0
                for cell in col:
                    if cell.row in (1, 2, footer_row_idx):
                        continue
                    val_str = str(cell.value or "")
                    if len(val_str) > max_len:
                        max_len = len(val_str)

                # Clamp width between 14 and 45 characters
                ws.column_dimensions[col_letter].width = max(14, min(max_len + 4, 45))

            # Save workbook to disk
            wb.save(out_path)
            _logger.info(f"BOM successfully exported to: {out_path}")
            return ExportResponse(status="success", file_path=out_path)

        except (IOError, OSError) as os_err:
            _logger.critical(f"Disk full or OS write error saving BOM to '{out_path}': {os_err}")
            if os.path.exists(out_path):
                try:
                    os.remove(out_path)
                except Exception:
                    pass
            return ExportResponse(
                status="error",
                file_path="",
                error_message=f"Cannot save generated document — disk error: {os_err}",
            )

        except Exception as render_err:
            _logger.error(f"Template rendering failed with openpyxl: {render_err}. Falling back to CSV.")
            return self._fallback_to_csv(request, out_path, str(render_err))

    def _fallback_to_csv(self, request: BomExportRequest, out_path: str, error_msg: str) -> ExportResponse:
        """
        Fallback to CSV export per §10.5 when Excel formatting fails.
        Maintains identical blank pricing columns and disclaimer banner.
        """
        csv_path = os.path.splitext(out_path)[0] + ".csv"
        try:
            with open(csv_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow([BOM_DISCLAIMER_TEXT])
                writer.writerow([BOM_COMPANY_HEADER])
                writer.writerow([f"Customer: {request.customer_name}", f"Query ID: {request.query_id}"])
                writer.writerow(BOM_HEADERS)

                for idx, item in enumerate(request.items, start=1):
                    cat_info = self._catalog.get(item.product_id, {})
                    oem = item.oem or cat_info.get("oem_name") or TBD_SALES_FALLBACK
                    pname = item.product_name or cat_info.get("product_name") or TBD_SALES_FALLBACK
                    pcat = item.product_category or cat_info.get("product_category") or TBD_SALES_FALLBACK
                    lic_model = item.licensing_model or cat_info.get("licensing_model") or TBD_SALES_FALLBACK
                    lic_unit = item.licensing_unit or cat_info.get("licensing_unit") or TBD_SALES_FALLBACK
                    deploy_model = item.deployment_model or cat_info.get("deployment_model") or TBD_SALES_FALLBACK

                    writer.writerow([
                        idx,
                        oem,
                        pname,
                        pcat,
                        item.product_id,
                        item.quantity,
                        lic_model,
                        item.license_term,
                        lic_unit,
                        deploy_model,
                        item.support_tier,
                        "",  # Unit Price BLANK
                        "",  # Total Price BLANK
                        item.remarks or "",
                    ])

            _logger.info(f"Fallback CSV BOM exported to: {csv_path}")
            return ExportResponse(
                status="success",
                file_path=csv_path,
                error_message=f"Fell back to CSV export due to formatting error: {error_msg}",
            )
        except Exception as csv_err:
            _logger.critical(f"Fallback CSV export also failed: {csv_err}")
            return ExportResponse(
                status="error",
                file_path="",
                error_message=f"BOM generation failed: {error_msg}. CSV fallback failed: {csv_err}",
            )


def export_bom(request: BomExportRequest) -> ExportResponse:
    """Helper function to export BOM using default BOMExporter instance."""
    exporter = BOMExporter()
    return exporter.export(request)


# Watermark & notice constants for BOQ
BOQ_WATERMARK_TEXT = "DRAFT — PRICING NOT INCLUDED"
BOQ_FOOTER_SALES_NOTE = "Pricing to be completed by iValue Sales Team"


class BOQExporter:
    """
    Generates formal Word proposal (.docx) with blank pricing columns per §8.4.
    """

    def __init__(self, catalog: Optional[Dict[str, Dict[str, str]]] = None) -> None:
        """
        catalog: Optional pre-loaded mapping of product_id -> product attributes.
        If omitted, loads automatically from dataset assets.
        """
        self._catalog: Dict[str, Dict[str, str]] = catalog if catalog is not None else {}
        if not self._catalog:
            self._load_catalog()

    def _load_catalog(self) -> None:
        """Loads product attributes from raw dataset or json assets."""
        dataset_path = os.path.join(
            os.path.dirname(__file__), "..", "..", "data", "raw", "iValue_Solution_Recommendation_Dataset.xlsx"
        )
        dataset_path = os.path.abspath(dataset_path)

        if os.path.isfile(dataset_path):
            try:
                wb = openpyxl.load_workbook(dataset_path, data_only=True, read_only=True)
                if "Products" in wb.sheetnames:
                    p_sheet = wb["Products"]
                    for row in p_sheet.iter_rows(min_row=2, values_only=True):
                        if row and row[0]:
                            pid = str(row[0]).strip()
                            self._catalog[pid] = {
                                "oem_name": str(row[2]).strip() if len(row) > 2 and row[2] else "",
                                "product_category": str(row[5]).strip() if len(row) > 5 and row[5] else "",
                                "product_name": str(row[6]).strip() if len(row) > 6 and row[6] else "",
                                "deployment_model": str(row[11]).strip() if len(row) > 11 and row[11] else "",
                                "licensing_model": "",
                                "licensing_unit": "",
                            }
                if "Product_Commercial" in wb.sheetnames:
                    c_sheet = wb["Product_Commercial"]
                    for row in c_sheet.iter_rows(min_row=2, values_only=True):
                        if row and row[0]:
                            pid = str(row[0]).strip()
                            if pid in self._catalog:
                                self._catalog[pid]["licensing_model"] = (
                                    str(row[3]).strip() if len(row) > 3 and row[3] else ""
                                )
                                self._catalog[pid]["licensing_unit"] = (
                                    str(row[5]).strip() if len(row) > 5 and row[5] else ""
                                )
                return
            except Exception as e:
                _logger.warning(f"Could not load Excel dataset catalog for BOQ: {e}")

        # Fallback to composite_products.json
        json_path = os.path.join(
            os.path.dirname(__file__), "..", "..", "data", "composite_products.json"
        )
        json_path = os.path.abspath(json_path)
        if os.path.isfile(json_path):
            try:
                import json
                with open(json_path, "r", encoding="utf-8") as f:
                    products = json.load(f)
                for item in products:
                    pid = item.get("product_id") or ""
                    if pid:
                        self._catalog[pid] = {
                            "oem_name": item.get("oem_name") or "",
                            "product_category": item.get("product_category") or "",
                            "product_name": item.get("product_name") or "",
                            "deployment_model": item.get("deployment_model") or "",
                            "licensing_model": item.get("licensing_model") or "",
                            "licensing_unit": item.get("licensing_unit") or "",
                        }
            except Exception as e:
                _logger.warning(f"Could not load json catalog for BOQ: {e}")

    def export(self, request: BoqExportRequest) -> ExportResponse:
        """
        Creates .docx via python-docx per §8.4 BOQ format:
        - Header: 'iValue InfoSolutions' (letterhead)
        - Metadata: Customer Name, Project Reference, Date, Prepared By, Query ID
        - Line items table mirroring BOM structure with blank price columns
        - Price column headers: 'Unit Price' and 'Total Price' (currency-neutral, optional currency_code)
        - All price cells: BLANK (strictly empty)
        - Footer note: 'Pricing to be completed by iValue Sales Team'
        - Watermark text in header: 'DRAFT — PRICING NOT INCLUDED'
        - Document version and generation timestamp in footer
        """
        if docx is None:
            _logger.error("python-docx library not installed. Falling back to CSV.")
            return self._fallback_to_csv(request, request.output_path, "python-docx not installed")

        out_path = os.path.abspath(request.output_path)
        out_dir = os.path.dirname(out_path)

        try:
            if out_dir:
                os.makedirs(out_dir, exist_ok=True)
        except OSError as e:
            _logger.critical(f"Disk or filesystem error creating directory '{out_dir}': {e}")
            return ExportResponse(
                status="error",
                file_path="",
                error_message=f"Cannot create directory '{out_dir}': {e}",
            )

        try:
            doc = docx.Document()

            # Configure Landscape Orientation for 14-column layout
            section = doc.sections[0]
            section.orientation = WD_ORIENT.LANDSCAPE
            section.page_width, section.page_height = section.page_height, section.page_width

            # 0.5-inch margins for spacious table fit
            section.top_margin = Inches(0.5)
            section.bottom_margin = Inches(0.5)
            section.left_margin = Inches(0.5)
            section.right_margin = Inches(0.5)

            # 1. Header with Watermark text (§8.4 / §12.4)
            header = section.header
            header_p = header.paragraphs[0]
            header_p.text = BOQ_WATERMARK_TEXT
            header_p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in header_p.runs:
                r.font.name = "Segoe UI"
                r.font.size = Pt(9)
                r.font.bold = True
                r.font.color.rgb = RGBColor(185, 28, 28)  # Crimson draft indicator

            # 2. Section Footer
            footer = section.footer
            footer_p = footer.paragraphs[0]
            footer_p.text = f"{BOQ_FOOTER_SALES_NOTE}  |  iValue PRISM v3.4  |  Confidential Draft"
            footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for r in footer_p.runs:
                r.font.name = "Segoe UI"
                r.font.size = Pt(8)
                r.font.color.rgb = RGBColor(148, 163, 184)

            # 3. Document Letterhead
            p_title = doc.add_paragraph()
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_title.paragraph_format.space_before = Pt(0)
            p_title.paragraph_format.space_after = Pt(2)
            r_title = p_title.add_run("iValue InfoSolutions")
            r_title.font.name = "Segoe UI"
            r_title.font.size = Pt(18)
            r_title.font.bold = True
            r_title.font.color.rgb = RGBColor(15, 23, 42)

            p_sub = doc.add_paragraph()
            p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_sub.paragraph_format.space_after = Pt(12)
            r_sub = p_sub.add_run("Bill of Quantities & Commercial Proposal Specification")
            r_sub.font.name = "Segoe UI"
            r_sub.font.size = Pt(11)
            r_sub.font.bold = True
            r_sub.font.color.rgb = RGBColor(71, 85, 105)

            # 4. Metadata Block
            p_meta = doc.add_paragraph()
            p_meta.paragraph_format.space_after = Pt(10)

            def _append_meta(p, label: str, val: str, is_last: bool = False) -> None:
                r_lbl = p.add_run(f"{label}: ")
                r_lbl.font.name = "Segoe UI"
                r_lbl.font.bold = True
                r_lbl.font.size = Pt(9.5)
                r_lbl.font.color.rgb = RGBColor(51, 65, 85)

                sep = "" if is_last else "    |    "
                r_val = p.add_run(f"{val}{sep}")
                r_val.font.name = "Segoe UI"
                r_val.font.size = Pt(9.5)
                r_val.font.color.rgb = RGBColor(15, 23, 42)

            _append_meta(p_meta, "Customer", request.customer_name or "Unspecified")
            _append_meta(p_meta, "Project Ref", request.project_reference or "N/A")
            _append_meta(p_meta, "Prepared By", request.prepared_by or "iValue Presales Engineering")
            _append_meta(p_meta, "Date", datetime.now().strftime("%Y-%m-%d"))
            _append_meta(p_meta, "Query ID", request.query_id or "N/A", is_last=True)

            # 5. Table Headers (14 columns)
            curr_tag = f" ({request.currency_code.strip()})" if request.currency_code and request.currency_code.strip() else ""
            unit_price_hdr = f"Unit Price{curr_tag}"
            total_price_hdr = f"Total Price{curr_tag}"

            headers = [
                "S.No.",
                "OEM",
                "Product Name",
                "Product Category",
                "Model/Part Reference",
                "Quantity",
                "Licensing Model",
                "License Term",
                "Licensing Unit",
                "Deployment Model",
                "Support Tier",
                unit_price_hdr,
                total_price_hdr,
                "Remarks",
            ]

            num_rows = 1 + len(request.items)
            table = doc.add_table(rows=num_rows, cols=14, style="Table Grid")

            # Header Row Formatting
            hdr_cells = table.rows[0].cells
            for col_i, h_text in enumerate(headers):
                hdr_cells[col_i].text = h_text
                # Background fill
                shd = OxmlElement("w:shd")
                shd.set(qn("w:val"), "clear")
                shd.set(qn("w:color"), "auto")
                shd.set(qn("w:fill"), "0F172A")
                hdr_cells[col_i]._tc.get_or_add_tcPr().append(shd)

                # Font & Alignment
                p_cell = hdr_cells[col_i].paragraphs[0]
                p_cell.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p_cell.runs:
                    run.font.name = "Segoe UI"
                    run.font.size = Pt(8.5)
                    run.font.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)

            # Data Rows Formatting
            for row_i, item in enumerate(request.items, start=1):
                row_cells = table.rows[row_i].cells
                cat_info = self._catalog.get(item.product_id, {})

                oem = item.oem or cat_info.get("oem_name") or TBD_SALES_FALLBACK
                pname = item.product_name or cat_info.get("product_name") or TBD_SALES_FALLBACK
                pcat = item.product_category or cat_info.get("product_category") or TBD_SALES_FALLBACK
                lic_model = item.licensing_model or cat_info.get("licensing_model") or TBD_SALES_FALLBACK
                lic_unit = item.licensing_unit or cat_info.get("licensing_unit") or TBD_SALES_FALLBACK
                deploy_model = item.deployment_model or cat_info.get("deployment_model") or TBD_SALES_FALLBACK

                row_values = [
                    str(row_i),
                    oem,
                    pname,
                    pcat,
                    item.product_id,
                    str(item.quantity),
                    lic_model,
                    item.license_term,
                    lic_unit,
                    deploy_model,
                    item.support_tier,
                    "",  # Col 11: Unit Price — ALWAYS BLANK
                    "",  # Col 12: Total Price — ALWAYS BLANK
                    item.remarks or "",
                ]

                # Alternating row tint
                fill_color = "F8FAFC" if (row_i % 2 == 0) else "FFFFFF"

                for col_i, val in enumerate(row_values):
                    row_cells[col_i].text = val
                    shd = OxmlElement("w:shd")
                    shd.set(qn("w:val"), "clear")
                    shd.set(qn("w:color"), "auto")
                    shd.set(qn("w:fill"), fill_color)
                    row_cells[col_i]._tc.get_or_add_tcPr().append(shd)

                    p_c = row_cells[col_i].paragraphs[0]
                    # Center align numeric and categorical codes
                    if col_i in (0, 4, 5, 7, 10, 11, 12):
                        p_c.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    else:
                        p_c.alignment = WD_ALIGN_PARAGRAPH.LEFT

                    for run in p_c.runs:
                        run.font.name = "Segoe UI"
                        run.font.size = Pt(8.5)
                        run.font.color.rgb = RGBColor(15, 23, 42)

            # 6. Body Footer Notice
            p_note = doc.add_paragraph()
            p_note.paragraph_format.space_before = Pt(14)
            p_note.paragraph_format.space_after = Pt(2)
            p_note.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_note = p_note.add_run(BOQ_FOOTER_SALES_NOTE)
            r_note.font.name = "Segoe UI"
            r_note.font.size = Pt(10.5)
            r_note.font.bold = True
            r_note.font.color.rgb = RGBColor(185, 28, 28)

            p_gen = doc.add_paragraph()
            p_gen.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_gen = p_gen.add_run(f"Document generated by iValue PRISM v3.4 — {datetime.now().isoformat()}")
            r_gen.font.name = "Segoe UI"
            r_gen.font.size = Pt(8.5)
            r_gen.font.color.rgb = RGBColor(148, 163, 184)

            doc.save(out_path)
            _logger.info(f"BOQ successfully exported to: {out_path}")
            return ExportResponse(status="success", file_path=out_path)

        except (IOError, OSError) as os_err:
            _logger.critical(f"Disk full or OS write error saving BOQ to '{out_path}': {os_err}")
            if os.path.exists(out_path):
                try:
                    os.remove(out_path)
                except Exception:
                    pass
            return ExportResponse(
                status="error",
                file_path="",
                error_message=f"Cannot save generated document — disk error: {os_err}",
            )

        except Exception as render_err:
            _logger.error(f"Template rendering failed with python-docx: {render_err}. Falling back to CSV.")
            return self._fallback_to_csv(request, out_path, str(render_err))

    def _fallback_to_csv(self, request: BoqExportRequest, out_path: str, error_msg: str) -> ExportResponse:
        """
        Fallback to CSV export per §10.5 when Word document generation fails.
        Maintains identical blank pricing columns and disclaimer banner.
        """
        csv_path = os.path.splitext(out_path)[0] + ".csv"
        try:
            with open(csv_path, "w", newline="", encoding="utf-8") as f:
                writer = csv.writer(f)
                writer.writerow([BOQ_WATERMARK_TEXT])
                writer.writerow(["iValue InfoSolutions — Bill of Quantities"])
                writer.writerow([
                    f"Customer: {request.customer_name}",
                    f"Project: {request.project_reference}",
                    f"Prepared By: {request.prepared_by}",
                    f"Query ID: {request.query_id}",
                ])

                curr_tag = f" ({request.currency_code.strip()})" if request.currency_code and request.currency_code.strip() else ""
                headers = [
                    "S.No.",
                    "OEM",
                    "Product Name",
                    "Product Category",
                    "Model/Part Reference",
                    "Quantity",
                    "Licensing Model",
                    "License Term",
                    "Licensing Unit",
                    "Deployment Model",
                    "Support Tier",
                    f"Unit Price{curr_tag}",
                    f"Total Price{curr_tag}",
                    "Remarks",
                ]
                writer.writerow(headers)

                for idx, item in enumerate(request.items, start=1):
                    cat_info = self._catalog.get(item.product_id, {})
                    oem = item.oem or cat_info.get("oem_name") or TBD_SALES_FALLBACK
                    pname = item.product_name or cat_info.get("product_name") or TBD_SALES_FALLBACK
                    pcat = item.product_category or cat_info.get("product_category") or TBD_SALES_FALLBACK
                    lic_model = item.licensing_model or cat_info.get("licensing_model") or TBD_SALES_FALLBACK
                    lic_unit = item.licensing_unit or cat_info.get("licensing_unit") or TBD_SALES_FALLBACK
                    deploy_model = item.deployment_model or cat_info.get("deployment_model") or TBD_SALES_FALLBACK

                    writer.writerow([
                        idx,
                        oem,
                        pname,
                        pcat,
                        item.product_id,
                        item.quantity,
                        lic_model,
                        item.license_term,
                        lic_unit,
                        deploy_model,
                        item.support_tier,
                        "",  # Unit Price BLANK
                        "",  # Total Price BLANK
                        item.remarks or "",
                    ])

            _logger.info(f"Fallback CSV BOQ exported to: {csv_path}")
            return ExportResponse(
                status="success",
                file_path=csv_path,
                error_message=f"Fell back to CSV export due to formatting error: {error_msg}",
            )
        except Exception as csv_err:
            _logger.critical(f"Fallback CSV export also failed: {csv_err}")
            return ExportResponse(
                status="error",
                file_path="",
                error_message=f"BOQ generation failed: {error_msg}. CSV fallback failed: {csv_err}",
            )


def export_boq(request: BoqExportRequest) -> ExportResponse:
    """Helper function to export BOQ using default BOQExporter instance."""
    exporter = BOQExporter()
    return exporter.export(request)
