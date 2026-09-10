"""
iValue PRISM — Document Ingestion and Text Extraction Module
SRS References: §3 Phase 1 step 3, §10.1, §10.2
Implementation Plan: TASK-P1.4

This module extracts plain text from requirement documents (.txt, .pdf, .docx),
performs strict size and format validation, detects scanned/image PDFs, catches
password-protected files, and handles character encoding fallbacks.
"""

from dataclasses import dataclass
import logging
import os
from typing import Optional, Tuple

import docx
import pdfplumber

try:
    from pdfminer.pdfdocument import PDFEncryptionError, PDFPasswordIncorrect
except ImportError:
    class PDFPasswordIncorrect(Exception):
        pass

    class PDFEncryptionError(Exception):
        pass

from src.utils.config import (
    MAX_FILE_SIZE_BYTES,
    SUPPORTED_EXTENSIONS,
    estimate_tokens,
)

_logger = logging.getLogger("prism.ingestion")


@dataclass
class ExtractionResult:
    """
    Structured outcome of parsing a requirement document (§3 Phase 1 step 3).
    """
    text: str
    source_filename: str
    file_size_kb: float
    char_count: int
    estimated_tokens: int
    warning: Optional[str] = None  # e.g., "scanned_pdf", "encoding_fallback"
    error: Optional[str] = None

    @property
    def success(self) -> bool:
        return self.error is None


class DocumentParser:
    """
    Extracts plain text from .txt, .pdf, and .docx files per SRS §3 and §10.2.
    """

    def extract(self, file_path: str) -> ExtractionResult:
        """
        Extracts plain text from a supported file, applying validation guards:
          - File existence
          - File size <= MAX_FILE_SIZE_BYTES (10 MB per §10.2)
          - Non-empty file (0 bytes rejected per NEG-09)
          - Supported extension in {'.txt', '.pdf', '.docx'} (NEG-08)
          - Password-protected PDF detection (NEG-11)
          - Scanned PDF detection: <50 chars text and >100 KB file size (NEG-12)
          - Character encoding fallback: UTF-8 -> latin-1 (NEG-14)
        """
        source_filename = os.path.basename(file_path)

        if not os.path.exists(file_path):
            return ExtractionResult(
                text="",
                source_filename=source_filename,
                file_size_kb=0.0,
                char_count=0,
                estimated_tokens=0,
                error=f"File not found: {file_path}",
            )

        try:
            file_size_bytes = os.path.getsize(file_path)
        except OSError as e:
            return ExtractionResult(
                text="",
                source_filename=source_filename,
                file_size_kb=0.0,
                char_count=0,
                estimated_tokens=0,
                error=f"Could not access file: {e}",
            )

        file_size_kb = round(file_size_bytes / 1024.0, 2)

        # 1. Zero-byte check (NEG-09)
        if file_size_bytes == 0:
            return ExtractionResult(
                text="",
                source_filename=source_filename,
                file_size_kb=0.0,
                char_count=0,
                estimated_tokens=0,
                error="The uploaded file is empty.",
            )

        # 2. Oversized file check (NEG-10)
        if file_size_bytes > MAX_FILE_SIZE_BYTES:
            max_mb = MAX_FILE_SIZE_BYTES // (1024 * 1024)
            return ExtractionResult(
                text="",
                source_filename=source_filename,
                file_size_kb=file_size_kb,
                char_count=0,
                estimated_tokens=0,
                error=f"File size exceeds the {max_mb} MB limit.",
            )

        # 3. Supported extension check (NEG-08)
        _, ext = os.path.splitext(file_path)
        ext_lower = ext.lower()
        if ext_lower not in SUPPORTED_EXTENSIONS:
            return ExtractionResult(
                text="",
                source_filename=source_filename,
                file_size_kb=file_size_kb,
                char_count=0,
                estimated_tokens=0,
                error=f"Unsupported file format '{ext}'. Please upload a .txt, .pdf, or .docx file.",
            )

        # 4. Dispatch to format-specific extractor
        extracted_text = ""
        warning: Optional[str] = None
        error: Optional[str] = None

        if ext_lower == ".txt":
            extracted_text, warning, error = self._extract_txt(file_path)
        elif ext_lower == ".pdf":
            extracted_text, warning, error = self._extract_pdf(file_path)
            # Scanned PDF detection (NEG-12): <50 chars text and file size > 100 KB
            if not error and len(extracted_text.strip()) < 50 and file_size_bytes > 100 * 1024:
                warning = "scanned_pdf"
        elif ext_lower == ".docx":
            extracted_text, warning, error = self._extract_docx(file_path)

        if error:
            return ExtractionResult(
                text="",
                source_filename=source_filename,
                file_size_kb=file_size_kb,
                char_count=0,
                estimated_tokens=0,
                warning=warning,
                error=error,
            )

        # Check for empty extracted content if not marked as scanned PDF
        if not extracted_text.strip() and warning != "scanned_pdf":
            return ExtractionResult(
                text="",
                source_filename=source_filename,
                file_size_kb=file_size_kb,
                char_count=0,
                estimated_tokens=0,
                warning=warning,
                error="The uploaded file contains no readable text.",
            )

        char_count = len(extracted_text)
        tokens = estimate_tokens(extracted_text)

        return ExtractionResult(
            text=extracted_text,
            source_filename=source_filename,
            file_size_kb=file_size_kb,
            char_count=char_count,
            estimated_tokens=tokens,
            warning=warning,
            error=None,
        )

    def _extract_txt(self, file_path: str) -> Tuple[str, Optional[str], Optional[str]]:
        """
        Reads plain text with UTF-8 encoding, falling back to latin-1 (NEG-14, §10.2).
        Returns: (text, warning, error)
        """
        try:
            with open(file_path, "r", encoding="utf-8") as f:
                return f.read(), None, None
        except UnicodeDecodeError:
            try:
                with open(file_path, "r", encoding="latin-1") as f:
                    _logger.warning(f"UTF-8 decoding failed for {file_path}; fell back to latin-1.")
                    return f.read(), "encoding_fallback", None
            except Exception as e:
                return "", None, f"Failed to read text file with fallback encoding: {e}"
        except Exception as e:
            return "", None, f"Failed to read text file: {e}"

    def _extract_pdf(self, file_path: str) -> Tuple[str, Optional[str], Optional[str]]:
        """
        Extracts text across all pages using pdfplumber (§10.2).
        Detects password protection (NEG-11).
        Returns: (text, warning, error)
        """
        try:
            with pdfplumber.open(file_path) as pdf:
                pages_text = []
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        pages_text.append(text)
                return "\n\n".join(pages_text), None, None
        except (PDFPasswordIncorrect, PDFEncryptionError) as e:
            _logger.warning(f"Password-protected PDF detected: {file_path} ({e})")
            return "", None, "The uploaded file is password-protected. Please remove password protection."
        except Exception as e:
            err_repr = repr(e).lower()
            err_str = str(e).lower()
            is_password = (
                "password" in err_repr
                or "encrypt" in err_repr
                or "password" in err_str
                or "encrypt" in err_str
                or any(
                    isinstance(arg, (PDFPasswordIncorrect, PDFEncryptionError))
                    or "password" in repr(arg).lower()
                    for arg in getattr(e, "args", [])
                )
            )
            if is_password:
                _logger.warning(f"Encrypted/password-protected PDF detected: {file_path} ({repr(e)})")
                return "", None, "The uploaded file is password-protected. Please remove password protection."
            _logger.error(f"Failed to parse PDF {file_path}: {e}")
            return "", None, f"Failed to parse PDF file: {e or repr(e)}"

    def _extract_docx(self, file_path: str) -> Tuple[str, Optional[str], Optional[str]]:
        """
        Extracts text from paragraphs, tables, and headers using python-docx (§10.2).
        Returns: (text, warning, error)
        """
        try:
            doc = docx.Document(file_path)
            content_blocks = []

            # 1. Headers
            for section in doc.sections:
                if section.header:
                    for p in section.header.paragraphs:
                        text = p.text.strip()
                        if text:
                            content_blocks.append(text)

            # 2. Main paragraphs
            for p in doc.paragraphs:
                text = p.text.strip()
                if text:
                    content_blocks.append(text)

            # 3. Table content (rows and cells)
            for table in doc.tables:
                for row in table.rows:
                    row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if row_cells:
                        content_blocks.append(" | ".join(row_cells))

            return "\n\n".join(content_blocks), None, None
        except Exception as e:
            _logger.error(f"Failed to parse DOCX {file_path}: {e}")
            return "", None, f"Failed to parse DOCX file: {e}"

    parse_file = extract
