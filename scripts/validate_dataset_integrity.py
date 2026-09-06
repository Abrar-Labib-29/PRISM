import sys
import codecs
import openpyxl
from collections import Counter

# Set standard output encoding to utf-8 safely
if sys.stdout.encoding != 'utf-8':
    try:
        sys.stdout.reconfigure(encoding='utf-8')
    except AttributeError:
        pass
import os

RAW_PATH = os.path.join('data', 'raw', 'iValue_Solution_Recommendation_Dataset.xlsx')
EXCEL_PATH = RAW_PATH if os.path.exists(RAW_PATH) else 'iValue_Solution_Recommendation_Dataset.xlsx'

EXPECTED_SHEETS = [
    'README',
    'Products',
    'Product_Features',
    'Product_Pros_Cons',
    'Product_Commercial',
    'Domain_Comparables',
    'Domain_Taxonomy'
]

CANONICAL_STATUS_VALUES = {'Confirmed', 'Draft', 'Internal-Only', 'TBD'}

CANONICAL_DOMAIN_CATEGORIES = {
    'Enterprise & Cyber Security',
    'Cloud & Application Life Management',
    'Infrastructure & Data Center',
    'Networking & Information Life Cycle Management'
}

def sheet_to_dicts(sheet):
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        return [], []
    headers = [str(h).strip() if h is not None else f"col_{i}" for i, h in enumerate(rows[0])]
    data = []
    for r in rows[1:]:
        if all(cell is None for cell in r):
            continue
        row_dict = {headers[i]: r[i] for i in range(min(len(headers), len(r)))}
        data.append(row_dict)
    return headers, data

def main():
    print("================================================================")
    print("       iValue Presales Automation - Phase 0.1 Validation        ")
    print("================================================================")
    print(f"Opening workbook: {EXCEL_PATH}...\n")

    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    all_passed = True

    # 1. Check Sheet Existence
    print("[TEST 1] Verifying Expected Sheets:")
    missing_sheets = set(EXPECTED_SHEETS) - set(wb.sheetnames)
    if missing_sheets:
        print(f"  [FAIL] Missing sheets: {missing_sheets}")
        all_passed = False
    else:
        print(f"  [PASS] All {len(EXPECTED_SHEETS)} expected sheets are present.")
        print(f"         Sheets: {wb.sheetnames}")

    # Load records
    _, prod_rows = sheet_to_dicts(wb['Products'])
    _, feat_rows = sheet_to_dicts(wb['Product_Features'])
    _, pros_rows = sheet_to_dicts(wb['Product_Pros_Cons'])
    _, comm_rows = sheet_to_dicts(wb['Product_Commercial'])
    _, comp_rows = sheet_to_dicts(wb['Domain_Comparables'])
    _, tax_rows = sheet_to_dicts(wb['Domain_Taxonomy'])

    # 2. Check Master Products
    print("\n[TEST 2] Verifying Products Master Record:")
    prod_ids = [r.get('Product_ID') for r in prod_rows if r.get('Product_ID')]
    unique_prod_ids = set(prod_ids)
    if len(prod_ids) != len(unique_prod_ids):
        duplicates = [item for item, count in Counter(prod_ids).items() if count > 1]
        print(f"  [FAIL] Duplicate Product_IDs in Products: {duplicates}")
        all_passed = False
    else:
        print(f"  [PASS] {len(unique_prod_ids)} unique products found. No duplicate primary keys.")

    # 3. Referential Integrity Check: Child Sheets -> Products
    print("\n[TEST 3] Referential Integrity (Child Sheets -> Products.Product_ID):")
    
    child_checks = [
        ('Product_Features', feat_rows, 'Product_ID'),
        ('Product_Pros_Cons', pros_rows, 'Product_ID'),
        ('Product_Commercial', comm_rows, 'Product_ID')
    ]

    for sheet_name, rows, col_name in child_checks:
        c_pids = {r.get(col_name) for r in rows if r.get(col_name)}
        orphaned = c_pids - unique_prod_ids
        missing_from_child = unique_prod_ids - c_pids
        
        if orphaned:
            print(f"  [FAIL] {sheet_name} has {len(orphaned)} orphaned {col_name}s: {orphaned}")
            all_passed = False
        else:
            print(f"  [PASS] {sheet_name} ({len(rows)} rows) -> 0 orphaned records.")

        if missing_from_child:
            print(f"         NOTE: {len(missing_from_child)} products have no entries in {sheet_name}.")
        else:
            print(f"         -> 100% product coverage in {sheet_name} ({len(c_pids)}/{len(unique_prod_ids)} products).")

    # 4. Referential Integrity: Domain_Comparables -> Products
    print("\n[TEST 4] Referential Integrity in Domain_Comparables:")
    comp_pids = set()
    for r in comp_rows:
        for prefix in ['Product_A', 'Product_B', 'Product_C', 'Product_D', 'Product_E']:
            pid = r.get(f'{prefix}_ID')
            if pid and str(pid).strip():
                comp_pids.add(str(pid).strip())

    comp_orphans = comp_pids - unique_prod_ids
    if comp_orphans:
        print(f"  [FAIL] Domain_Comparables has {len(comp_orphans)} orphaned Product_IDs: {comp_orphans}")
        all_passed = False
    else:
        print(f"  [PASS] Domain_Comparables ({len(comp_rows)} rows) references {len(comp_pids)} products -> 0 orphaned references.")

    # 5. Taxonomy Referential Integrity
    print("\n[TEST 5] Taxonomy Integrity & Alignment:")
    tax_subdomains = {r.get('Sub_Domain') for r in tax_rows if r.get('Sub_Domain')}
    print(f"  Total defined sub-domains in Domain_Taxonomy: {len(tax_subdomains)}")

    # Check Products.Sub_Domain
    prod_subdomains = {r.get('Sub_Domain') for r in prod_rows if r.get('Sub_Domain')}
    prod_tax_mismatch = prod_subdomains - tax_subdomains
    if prod_tax_mismatch:
        print(f"  [FAIL] Products has Sub_Domains not in Domain_Taxonomy: {prod_tax_mismatch}")
        all_passed = False
    else:
        print(f"  [PASS] All {len(prod_subdomains)} Sub_Domains in Products match Domain_Taxonomy.")

    # Check Domain_Comparables.Domain_Category
    comp_subdomains = {r.get('Domain_Category') for r in comp_rows if r.get('Domain_Category')}
    comp_tax_mismatch = comp_subdomains - tax_subdomains
    if comp_tax_mismatch:
        print(f"  [FAIL] Domain_Comparables has categories not in Domain_Taxonomy: {comp_tax_mismatch}")
        all_passed = False
    else:
        print(f"  [PASS] All {len(comp_subdomains)} categories in Domain_Comparables match Domain_Taxonomy.")

    # Check Domain_Category valid values
    prod_domain_cats = {r.get('Domain_Category') for r in prod_rows if r.get('Domain_Category')}
    invalid_cats = prod_domain_cats - CANONICAL_DOMAIN_CATEGORIES
    if invalid_cats:
        print(f"  [FAIL] Products has invalid Domain_Category values: {invalid_cats}")
        all_passed = False
    else:
        print(f"  [PASS] All Products belong to one of the 4 canonical Domain_Categories.")

    # 6. Canonical Status Values Check
    print("\n[TEST 6] Data_Status Standardization Check:")
    for sheet_name, rows in [
        ('Products', prod_rows),
        ('Product_Commercial', comm_rows),
        ('Domain_Taxonomy', tax_rows)
    ]:
        statuses = set(r.get('Data_Status') for r in rows if r.get('Data_Status'))
        invalid_statuses = statuses - CANONICAL_STATUS_VALUES
        if invalid_statuses:
            print(f"  [FAIL] {sheet_name} contains non-canonical Data_Status: {invalid_statuses}")
            all_passed = False
        else:
            print(f"  [PASS] {sheet_name} uses only canonical Data_Status values: {statuses}")

    print("\n================================================================")
    if all_passed:
        print("[SUCCESS] ALL INTEGRITY TESTS PASSED: 0 orphaned rows, 100% consistent.")
        print("================================================================")
        sys.exit(0)
    else:
        print("[FAILURE] INTEGRITY VERIFICATION FAILED. Review errors above.")
        print("================================================================")
        sys.exit(1)

if __name__ == "__main__":
    main()
