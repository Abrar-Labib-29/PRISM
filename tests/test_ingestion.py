"""
Unit tests for DocumentParser ingestion (.txt, .docx, error handling).
SRS References: §3 Phase 1 step 3, §10.1, §10.2
Implementation Plan: TASK-P1.4
"""

import os
import tempfile
import unittest

import docx

from src.core.ingestion import DocumentParser, ExtractionResult


class TestIngestion(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.parser = DocumentParser()

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_text_file_extraction(self):
        txt_path = os.path.join(self.temp_dir.name, "requirement.txt")
        sample_content = "Customer requires high-throughput next-generation firewall with 10 Gbps SSL inspection."
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(sample_content)

        result: ExtractionResult = self.parser.extract(txt_path)
        self.assertTrue(result.success)
        self.assertIsNone(result.error)
        self.assertEqual(result.text, sample_content)
        self.assertEqual(result.char_count, len(sample_content))
        self.assertGreater(result.estimated_tokens, 0)
        self.assertGreater(result.file_size_kb, 0.0)

    def test_docx_file_extraction(self):
        docx_path = os.path.join(self.temp_dir.name, "requirement.docx")
        doc = docx.Document()
        doc.add_heading("Customer Requirement Document", level=1)
        doc.add_paragraph("Need PAM solution with credential vaulting and session isolation.")
        doc.save(docx_path)

        result: ExtractionResult = self.parser.extract(docx_path)
        self.assertTrue(result.success)
        self.assertIsNone(result.error)
        self.assertIn("Customer Requirement Document", result.text)
        self.assertIn("Need PAM solution", result.text)

    def test_zero_byte_file_rejected(self):
        empty_path = os.path.join(self.temp_dir.name, "empty.txt")
        with open(empty_path, "w", encoding="utf-8") as f:
            pass  # 0 bytes

        result: ExtractionResult = self.parser.extract(empty_path)
        self.assertFalse(result.success)
        self.assertIsNotNone(result.error)
        self.assertIn("empty", result.error.lower())

    def test_unsupported_file_extension(self):
        unsupported = os.path.join(self.temp_dir.name, "data.csv")
        with open(unsupported, "w", encoding="utf-8") as f:
            f.write("col1,col2\nval1,val2")

        result: ExtractionResult = self.parser.extract(unsupported)
        self.assertFalse(result.success)
        self.assertIsNotNone(result.error)
        self.assertIn("unsupported", result.error.lower())

    def test_nonexistent_file(self):
        result = self.parser.extract("C:/nonexistent_prism_file.txt")
        self.assertFalse(result.success)
        self.assertIsNotNone(result.error)

    def test_encoding_fallback_latin1(self):
        txt_path = os.path.join(self.temp_dir.name, "latin1.txt")
        # Write bytes that are valid latin-1 but invalid utf-8
        with open(txt_path, "wb") as f:
            f.write(b"Technical requirement with \xe9\xe8\xe0 accented chars")

        result = self.parser.extract(txt_path)
        self.assertTrue(result.success)
        self.assertIn("Technical requirement", result.text)


if __name__ == "__main__":
    unittest.main()
