"""
Unit tests for BOM (.xlsx) and BOQ (.docx) proposal document exporters.
SRS References: §3 Phase 2 step 3, §8.3, §8.4, §10.5, §9.10
Implementation Plan: TASK-P1.9
"""

import os
import tempfile
import unittest

import openpyxl
import docx

from src.core.exporters import (
    BOM_DISCLAIMER_TEXT,
    BOM_HEADERS,
    BOMExporter,
    BOQExporter,
    BomExportRequest,
    BomItem,
    BoqExportRequest,
    ExportResponse,
)


class TestExporters(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.sample_items = [
            BomItem(
                product_id="OEM-001-P01",
                quantity=2,
                license_term="3-Year",
                support_tier="24x7 Enterprise",
                oem="Fortinet",
                product_name="FortiGate 100F",
                product_category="Next-Generation Firewall",
                licensing_model="Subscription",
                deployment_model="Hardware Appliance",
                remarks="High availability pair",
            ),
            BomItem(
                product_id="OEM-002-P01",
                quantity=300,
                license_term="1-Year",
                support_tier="Standard",
                oem="CyberArk",
                product_name="Privilege Cloud",
                product_category="Privileged Access Management",
                licensing_model="Per User Subscription",
                deployment_model="SaaS",
            ),
        ]

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_bom_export_excel_structure(self):
        out_path = os.path.join(self.temp_dir.name, "test_bom.xlsx")
        req = BomExportRequest(
            query_id="req-10001",
            customer_name="Acme Corp",
            items=self.sample_items,
            output_path=out_path,
        )

        exporter = BOMExporter()
        resp: ExportResponse = exporter.export(req)
        self.assertEqual(resp.status, "success")
        self.assertTrue(os.path.exists(out_path))

        # Open workbook and inspect
        wb = openpyxl.load_workbook(out_path)
        ws = wb.active

        # Find header row
        header_row = None
        for row in ws.iter_rows(values_only=True):
            if row and row[0] == "S.No.":
                header_row = [str(c) for c in row if c is not None]
                break

        self.assertIsNotNone(header_row)
        for h in BOM_HEADERS:
            self.assertIn(h, header_row)

        # STRICT ZERO-PRICING CHECK: Ensure Price columns are completely empty
        unit_price_col_idx = header_row.index("Unit Price — To be filled by Sales") + 1
        total_price_col_idx = header_row.index("Total Price — To be filled by Sales") + 1

        for r in range(6, 6 + len(self.sample_items)):
            up_val = ws.cell(row=r, column=unit_price_col_idx).value
            tp_val = ws.cell(row=r, column=total_price_col_idx).value
            # Must be None, blank, or placeholder disclaimer
            self.assertTrue(up_val is None or up_val == "" or "Sales" in str(up_val))
            self.assertTrue(tp_val is None or tp_val == "" or "Sales" in str(tp_val))

    def test_boq_export_word_document(self):
        out_path = os.path.join(self.temp_dir.name, "test_boq.docx")
        req = BoqExportRequest(
            query_id="req-20002",
            customer_name="Global Enterprise Ltd",
            project_reference="Cybersecurity Refresh 2026",
            prepared_by="iValue Presales Team",
            items=self.sample_items,
            output_path=out_path,
        )

        exporter = BOQExporter()
        resp: ExportResponse = exporter.export(req)
        self.assertEqual(resp.status, "success")
        self.assertTrue(os.path.exists(out_path))

        # Inspect DOCX
        doc = docx.Document(out_path)
        full_text = "\n".join(p.text for p in doc.paragraphs)

        self.assertIn("Bill of Quantities", full_text)
        self.assertIn("Global Enterprise Ltd", full_text)
        self.assertIn("Cybersecurity Refresh 2026", full_text)

        # Ensure disclaimer is present
        self.assertIn("pricing", full_text.lower())
        self.assertIn("sales", full_text.lower())


if __name__ == "__main__":
    unittest.main()
